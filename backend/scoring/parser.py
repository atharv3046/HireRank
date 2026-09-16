"""
Standalone Resume & Document Parser
===================================
Extracts clean, normalized text from PDF and DOCX files.
Segments text into canonical resume sections.
Flags unparseable or corrupt files with needs_manual_review=True.
Zero FastAPI or DB imports.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional, Dict, List, Tuple
import logging

logger = logging.getLogger(__name__)


@dataclass
class ResumeSection:
    """A labelled segment of a resume."""
    label: str          # 'contact' | 'education' | 'experience' | 'skills' | 'projects' | 'summary' | 'other'
    raw_text: str       # content within this section


@dataclass
class ParsedResume:
    """Standardized output of ResumeParser."""
    file_path: str
    raw_text: str                          # full normalized text
    sections: Dict[str, str]              # label -> content string
    needs_manual_review: bool = False     # true if file failed to parse or is empty
    parse_status: str = "ok"              # 'ok' | 'needs_manual_review' | 'failed'
    parse_notes: List[str] = field(default_factory=list)

    @property
    def contact_text(self) -> str:
        return self.sections.get("contact", "")

    @property
    def education_text(self) -> str:
        return self.sections.get("education", "")

    @property
    def experience_text(self) -> str:
        return self.sections.get("experience", "")

    @property
    def skills_text(self) -> str:
        return self.sections.get("skills", "")

    @property
    def projects_text(self) -> str:
        return self.sections.get("projects", "")

    @property
    def summary_text(self) -> str:
        return self.sections.get("summary", "")


# ──────────────────────────────────────────────────────────────────────────────
# Section Patterns
# ──────────────────────────────────────────────────────────────────────────────

SECTION_PATTERNS: List[Tuple[str, List[str], int]] = [
    ("contact", [
        r"contact\s*(information|details|info)?",
        r"personal\s*(information|details|profile)?",
        r"(phone|email|address|linkedin)\s*:",
    ], 1),

    ("education", [
        r"education(al)?\s*(background|qualifications|history)?",
        r"academic\s*(background|record|history|qualifications)",
        r"qualifications?",
        r"degrees?\s*(and\s*certifications?)?",
    ], 2),

    ("experience", [
        r"(work|professional|employment|career|job)\s*(experience|history|background)?",
        r"experience",
        r"internship(s)?",
    ], 3),

    ("skills", [
        r"(technical\s+|core\s+|key\s+)?skills?",
        r"technologies\s*(and\s*tools)?",
        r"tech\s*stack",
        r"competenc(ies|e)",
        r"proficienc(ies|y)",
        r"areas\s+of\s+expertise",
    ], 4),

    ("projects", [
        r"(personal\s+|key\s+|academic\s+)?projects?",
        r"open\s*source\s*contributions?",
    ], 5),

    ("summary", [
        r"(professional\s+|executive\s+|career\s+)?summary",
        r"about\s*me",
        r"objective",
    ], 6),

    ("certifications", [
        r"certificat(ions?|es?)",
        r"licenses?\s*(and\s*certifications?)?",
    ], 7),

    ("publications", [
        r"publications?",
        r"research\s*papers?",
    ], 8),
]

# Character mappings
UNICODE_REPLACEMENTS = {
    "\u2018": "'", "\u2019": "'",
    "\u201c": '"', "\u201d": '"',
    "\u2013": "-", "\u2014": "-",
    "\u2022": "\n- ", "\u25e6": "\n- ", "\u25aa": "\n- ", "\u2023": "\n- ",
    "\u00a0": " ",  # non-breaking space
    "\uf0b7": "\n- ",  # common Wingdings bullet
    "ﬁ": "fi", "ﬂ": "fl", "ﬀ": "ff", "ﬃ": "ffi", "ﬄ": "ffl",
}


def normalise_text(raw: str) -> str:
    """Normalize unicode, bullets, ligatures, and extra whitespace."""
    if not raw:
        return ""

    text = raw
    for src, dst in UNICODE_REPLACEMENTS.items():
        text = text.replace(src, dst)

    text = unicodedata.normalize("NFKD", text)

    # Clean non-printable characters keeping tabs & newlines
    text = re.sub(r"[^\x09\x0A\x0D\x20-\x7E]", " ", text)

    # Collapse horizontal spaces
    text = re.sub(r"[ \t]+", " ", text)

    # Collapse excessive blank lines
    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()


def is_heading_line(line: str) -> Tuple[Optional[str], int]:
    """Check if line is a section heading."""
    stripped = line.strip().rstrip(":-_ ").strip()
    if not stripped:
        return None, 999

    # Headings are concise
    if len(stripped.split()) > 6 or len(stripped) > 45:
        return None, 999

    stripped_lower = stripped.lower()

    for label, patterns, priority in SECTION_PATTERNS:
        for pat in patterns:
            if re.fullmatch(pat, stripped_lower, re.IGNORECASE):
                return label, priority

    return None, 999


def segment_sections(text: str) -> Dict[str, str]:
    """Segment text into canonical sections."""
    lines = text.splitlines()
    sections: Dict[str, List[str]] = {}
    current_section = "other"
    sections[current_section] = []

    for line in lines:
        label, _ = is_heading_line(line)
        if label:
            current_section = label
            if current_section not in sections:
                sections[current_section] = []
        else:
            sections[current_section].append(line)

    result: Dict[str, str] = {}
    for label, sec_lines in sections.items():
        content = "\n".join(sec_lines).strip()
        if content:
            result[label] = content

    return result


class ResumeParser:
    """
    Parser for PDF, DOCX, and text resumes with fallback extraction
    and needs_manual_review handling.
    """

    @staticmethod
    def extract_pdf(file_path: Path) -> Tuple[str, List[str]]:
        notes: List[str] = []
        extracted_text = ""

        # Primary: pdfplumber
        try:
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                pages_text = []
                for idx, page in enumerate(pdf.pages):
                    pt = page.extract_text() or ""
                    if pt.strip():
                        pages_text.append(pt)
                extracted_text = "\n".join(pages_text)
        except Exception as e:
            notes.append(f"pdfplumber extraction failed: {e}")

        # Fallback: pypdf
        if not extracted_text.strip():
            try:
                import pypdf
                reader = pypdf.PdfReader(str(file_path))
                pages_text = []
                for page in reader.pages:
                    pt = page.extract_text() or ""
                    if pt.strip():
                        pages_text.append(pt)
                extracted_text = "\n".join(pages_text)
                if extracted_text.strip():
                    notes.append("Recovered via pypdf fallback")
            except Exception as e:
                notes.append(f"pypdf fallback failed: {e}")

        return extracted_text, notes

    @staticmethod
    def extract_docx(file_path: Path) -> Tuple[str, List[str]]:
        notes: List[str] = []
        try:
            import docx
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            # Include tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            paragraphs.append(cell.text.strip())
            return "\n".join(paragraphs), notes
        except Exception as e:
            notes.append(f"python-docx extraction failed: {e}")
            return "", notes

    @classmethod
    def parse(cls, file_path: str | Path) -> ParsedResume:
        p = Path(file_path)
        if not p.exists():
            return ParsedResume(
                file_path=str(p),
                raw_text="",
                sections={},
                needs_manual_review=True,
                parse_status="failed",
                parse_notes=[f"File not found: {file_path}"]
            )

        suffix = p.suffix.lower()
        raw_text = ""
        notes: List[str] = []

        if suffix == ".pdf":
            raw_text, notes = cls.extract_pdf(p)
        elif suffix in (".docx", ".doc"):
            raw_text, notes = cls.extract_docx(p)
        elif suffix in (".txt", ".md"):
            try:
                raw_text = p.read_text(encoding="utf-8", errors="replace")
            except Exception as e:
                notes.append(f"Text reading failed: {e}")
        else:
            return ParsedResume(
                file_path=str(p),
                raw_text="",
                sections={},
                needs_manual_review=True,
                parse_status="needs_manual_review",
                parse_notes=[f"Unsupported file format: {suffix}"]
            )

        normalised = normalise_text(raw_text)

        # Check if unparseable or empty
        if len(normalised) < 25:
            notes.append("Extracted text is empty or too short (< 25 characters). Possible scanned/image PDF.")
            return ParsedResume(
                file_path=str(p),
                raw_text=normalised,
                sections={"other": normalised} if normalised else {},
                needs_manual_review=True,
                parse_status="needs_manual_review",
                parse_notes=notes
            )

        sections = segment_sections(normalised)

        return ParsedResume(
            file_path=str(p),
            raw_text=normalised,
            sections=sections,
            needs_manual_review=False,
            parse_status="ok",
            parse_notes=notes
        )
