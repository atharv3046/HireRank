"""
Tests for Phase 2 Anonymous Guest Screening Endpoints
=====================================================
Validates POST /guest/screen, GET /guest/session/{id}/status,
and GET /guest/session/{id}/results (and alias /preview)
wired to the standalone scoring package.
"""

import io
import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.routers.guest import _SESSIONS

client = TestClient(app)

SAMPLE_JD = """
Senior Python Developer
We are looking for an experienced Senior Python Developer with strong backend skills.
Requirements:
- 4+ years of Python development experience
- Experience with FastAPI, Docker, and PostgreSQL
- Bachelor's degree in Computer Science or related engineering field
"""

SAMPLE_RESUME_TEXT = """
Jane Doe
jane.doe@example.com
(555) 123-4567

EXPERIENCE
Senior Software Engineer - Acme Cloud (Jan 2020 - Present)
- Developed high-throughput REST APIs using Python and FastAPI
- Containerized microservices using Docker and deployed to Kubernetes
- Optimized SQL queries and data models in PostgreSQL

EDUCATION
B.S. in Computer Science - Tech University (2016 - 2020)

SKILLS
Python, FastAPI, Docker, PostgreSQL, Redis, Git
"""


class TestPhase2Endpoints:

    def test_screen_validation_short_jd(self):
        response = client.post(
            "/guest/screen",
            data={"job_description": "Short JD"},
            files=[("files", ("resume.pdf", b"%PDF-1.4 dummy", "application/pdf"))]
        )
        assert response.status_code == 422
        assert "at least 50 characters" in response.json()["detail"]

    def test_screen_validation_no_files(self):
        response = client.post(
            "/guest/screen",
            data={"job_description": SAMPLE_JD},
        )
        assert response.status_code == 422

    def test_screen_validation_unsupported_file_extension(self):
        response = client.post(
            "/guest/screen",
            data={"job_description": SAMPLE_JD},
            files=[("files", ("resume.exe", b"binary", "application/octet-stream"))]
        )
        assert response.status_code == 422
        assert "unsupported format" in response.json()["detail"]

    def test_guest_screen_and_polling_flow(self):
        # Generate a valid in-memory DOCX resume to test complete extraction and scoring
        import docx
        doc = docx.Document()
        doc.add_heading("Jane Doe", level=1)
        doc.add_paragraph("jane.doe@example.com | (555) 123-4567")
        doc.add_heading("Experience", level=2)
        doc.add_paragraph("Senior Software Engineer at Acme Cloud (Jan 2020 - Present)")
        doc.add_paragraph("Built microservices with Python, FastAPI, Docker, and PostgreSQL")
        doc.add_heading("Education", level=2)
        doc.add_paragraph("B.S. in Computer Science (2016 - 2020)")
        doc.add_heading("Skills", level=2)
        doc.add_paragraph("Python, FastAPI, Docker, PostgreSQL, Redis")
        docx_buf = io.BytesIO()
        doc.save(docx_buf)
        docx_bytes = docx_buf.getvalue()

        # Submit valid screening
        files = [
            ("files", ("jane_doe.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))
        ]
        res = client.post(
            "/guest/screen",
            data={"job_description": SAMPLE_JD},
            files=files
        )
        assert res.status_code == 200
        data = res.json()
        assert "session_id" in data
        assert "job_id" in data
        assert data["candidate_count"] == 1

        session_id = data["session_id"]

        # Check status endpoint
        status_res = client.get(f"/guest/session/{session_id}/status")
        assert status_res.status_code == 200
        status_data = status_res.json()
        assert status_data["session_id"] == session_id
        assert status_data["status"] in ("processing", "done")
        assert "stage" in status_data

        # Wait or manually trigger pipeline if synchronous for test
        from app.routers.guest import _run_guest_pipeline
        sess = _SESSIONS.get(session_id)
        if sess and sess["status"] != "done":
            _run_guest_pipeline(session_id, sess["job_id"], sess["candidate_ids"])

        # Fetch results
        results_res = client.get(f"/guest/session/{session_id}/results")
        assert results_res.status_code == 200
        res_data = results_res.json()

        assert res_data["session_id"] == session_id
        assert len(res_data["candidates"]) == 1
        cand = res_data["candidates"][0]

        # Verify masking
        assert cand["name_masked"] == "Candidate #1"
        assert "@" in cand["email_masked"]
        assert "****" in cand["email_masked"]

        # Verify score & tier
        assert cand["score"] is not None
        assert cand["score"] > 50.0
        assert cand["tier"] in ("Strong", "Potential")
        assert "Python" in cand["matched_skills"]

        # Verify shared aggregate stats
        stats = res_data["stats"]
        assert stats["total_processed"] == 1
        assert stats["average_score"] == cand["score"]
        assert "tier_distribution" in stats

        # Verify preview alias produces same result
        preview_res = client.get(f"/guest/session/{session_id}/preview")
        assert preview_res.status_code == 200
        assert preview_res.json() == res_data

    def test_unparseable_corrupt_file_handled_gracefully(self):
        # Submit corrupted file
        files = [
            ("files", ("corrupt.pdf", b"%PDF-1.4 unreadable binary garbage without catalog", "application/pdf"))
        ]
        res = client.post(
            "/guest/screen",
            data={"job_description": SAMPLE_JD},
            files=files
        )
        assert res.status_code == 200
        session_id = res.json()["session_id"]

        from app.routers.guest import _run_guest_pipeline
        sess = _SESSIONS.get(session_id)
        _run_guest_pipeline(session_id, sess["job_id"], sess["candidate_ids"])

        results_res = client.get(f"/guest/session/{session_id}/results")
        assert results_res.status_code == 200
        data = results_res.json()

        cand = data["candidates"][0]
        # Candidate should have needs_manual_review=True, null score, and Needs Review tier
        assert cand["needs_manual_review"] is True
        assert cand["score"] is None
        assert cand["tier"] == "Needs Review"
        assert data["stats"]["needs_review_count"] == 1
